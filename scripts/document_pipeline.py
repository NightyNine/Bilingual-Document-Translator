#!/usr/bin/env python3
"""Checkpointed bilingual DOCX/PDF translation pipeline.

The script is deliberately model-agnostic. Hermes supplies JSON responses for
analysis, translation, and review batches; this program owns extraction,
validation, OOXML insertion, rendering, and resumability.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import zipfile
from copy import deepcopy
from pathlib import Path
from typing import Any, Iterable

try:
    from lxml import etree
except ImportError:  # pragma: no cover - dependency check handles this
    etree = None

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = {"w": W_NS, "r": R_NS}
W = f"{{{W_NS}}}"
STORY_RE = re.compile(r"^word/(document|header\d+|footer\d+|footnotes|endnotes|comments)\.xml$")
SKIP_TEXT_RE = re.compile(r"^(?:https?://|www\.|[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}$)")
ENGLISH_WORD_RE = re.compile(r"[A-Za-z][A-Za-z-]{2,}")
VISIBLE_CJK_FONT = "Kaiti SC"
VISIBLE_LATIN_FONT = "Cambria"


def die(message: str, code: int = 2) -> None:
    print(json.dumps({"ok": False, "error": message}, ensure_ascii=False), file=sys.stderr)
    raise SystemExit(code)


def read_json(path: Path, default: Any = None) -> Any:
    if not path.exists():
        return default
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def sha256_text(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def normalize_text(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\u00a0", " ")).strip()


def cjk_count(value: str) -> int:
    return sum("CJK UNIFIED" in unicodedata.name(ch, "") for ch in value)


def latin_count(value: str) -> int:
    return sum("LATIN" in unicodedata.name(ch, "") for ch in value)


def detect_direction(text: str) -> str:
    """Return zh-to-en, en-to-zh, or keep for non-linguistic content."""
    cleaned = normalize_text(text)
    if not cleaned or SKIP_TEXT_RE.match(cleaned):
        return "keep"
    zh, en = cjk_count(cleaned), latin_count(cleaned)
    # A paragraph that already contains a substantive Chinese title followed
    # by its English equivalent is already bilingual. Treat it as complete so
    # the finalizer does not append a duplicate Chinese or English paragraph.
    if zh >= 2 and en >= 8 and len(ENGLISH_WORD_RE.findall(cleaned)) >= 2:
        return "keep"
    if zh == 0 and en == 0:
        return "keep"
    if zh >= max(2, en * 0.2):
        return "zh-to-en"
    if en >= max(2, zh * 0.2):
        return "en-to-zh"
    return "keep"


def require_lxml() -> None:
    if etree is None:
        die("Missing dependency: lxml. Install scripts/requirements.txt before processing DOCX files.")


def list_story_parts(archive: zipfile.ZipFile) -> list[str]:
    return sorted(name for name in archive.namelist() if STORY_RE.match(name))


def paragraph_text(p: Any) -> str:
    pieces: list[str] = []
    for node in p.iter():
        if node.tag == f"{W}t" or node.tag == f"{W}delText":
            pieces.append(node.text or "")
        elif node.tag == f"{W}tab":
            pieces.append("\t")
        elif node.tag in {f"{W}br", f"{W}cr"}:
            pieces.append("\n")
    return "".join(pieces)


def kind_for_part(part: str) -> str:
    name = Path(part).name
    if name == "document.xml":
        return "body"
    if name.startswith("header"):
        return "header"
    if name.startswith("footer"):
        return "footer"
    if name == "footnotes.xml":
        return "footnote"
    if name == "endnotes.xml":
        return "endnote"
    return "comment"


def extract_docx_units(docx_path: Path) -> list[dict[str, Any]]:
    require_lxml()
    units: list[dict[str, Any]] = []
    with zipfile.ZipFile(docx_path) as archive:
        for part in list_story_parts(archive):
            root = etree.fromstring(archive.read(part))
            paragraphs = root.xpath(".//w:p", namespaces=NS)
            kind = kind_for_part(part)
            for index, p in enumerate(paragraphs):
                text = paragraph_text(p)
                clean = normalize_text(text)
                if not clean:
                    continue
                direction = detect_direction(clean)
                units.append(
                    {
                        "id": f"{part}:{index}",
                        "part": part,
                        "paragraph_index": index,
                        "kind": kind,
                        "text": text,
                        "normalized_text": clean,
                        "source_hash": sha256_text(text),
                        "direction": direction,
                        "translatable": direction != "keep",
                    }
                )
    return units


def pdf_text_pages(pdf_path: Path) -> tuple[list[str], bool, list[str]]:
    try:
        import fitz  # type: ignore
    except ImportError:
        die("Missing dependency: PyMuPDF. Install scripts/requirements.txt before processing PDF files.")
    doc = fitz.open(pdf_path)
    pages: list[str] = []
    warnings: list[str] = []
    scanned = False
    for page_no, page in enumerate(doc, 1):
        text = page.get_text("text") or ""
        if len(normalize_text(text)) < 20:
            scanned = True
            try:
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                from PIL import Image  # type: ignore
                import io
                image = Image.open(io.BytesIO(pix.tobytes("png")))
                import pytesseract  # type: ignore
                ocr = pytesseract.image_to_string(image, lang="eng+chi_sim")
                if normalize_text(ocr):
                    text = ocr
                else:
                    warnings.append(f"Page {page_no}: OCR returned no text")
            except Exception as exc:
                warnings.append(f"Page {page_no}: OCR failed ({exc})")
        pages.append(text)
    return pages, scanned, warnings


def make_docx_from_pdf(pdf_path: Path, output_docx: Path) -> tuple[bool, list[str]]:
    """Create a reflowable intermediate DOCX from PDF pages."""
    try:
        from docx import Document  # type: ignore
        from docx.enum.text import WD_BREAK  # type: ignore
    except ImportError:
        die("Missing dependency: python-docx. Install scripts/requirements.txt before processing PDF files.")
    pages, scanned, warnings = pdf_text_pages(pdf_path)
    document = Document()
    for page_no, page_text in enumerate(pages):
        chunks = [chunk.strip() for chunk in re.split(r"\n\s*\n+", page_text) if chunk.strip()]
        if not chunks and normalize_text(page_text):
            chunks = [normalize_text(page_text)]
        for chunk in chunks:
            document.add_paragraph(chunk)
        if page_no < len(pages) - 1:
            p = document.add_paragraph()
            p.add_run().add_break(WD_BREAK.PAGE)
    output_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_docx)
    if scanned:
        warnings.insert(0, "Scanned PDF: OCR/reflow output is approximate; inspect flagged pages visually.")
    return scanned, warnings


def prepare(args: argparse.Namespace) -> None:
    input_path = Path(args.input).expanduser().resolve()
    if not input_path.exists() or not input_path.is_file():
        die(f"Input file does not exist: {input_path}")
    work = Path(args.work_dir).expanduser().resolve()
    work.mkdir(parents=True, exist_ok=True)
    ext = input_path.suffix.lower()
    if ext not in {".docx", ".pdf"}:
        die("Only .docx and .pdf inputs are supported.")
    source_copy = work / f"source{ext}"
    shutil.copy2(input_path, source_copy)
    warnings: list[str] = []
    scanned = False
    if ext == ".pdf":
        intermediate = work / "source.docx"
        scanned, warnings = make_docx_from_pdf(source_copy, intermediate)
        docx_path = intermediate
    else:
        docx_path = source_copy
    units = extract_docx_units(docx_path)
    manifest = {
        "version": 1,
        "input": str(input_path),
        "input_sha256": sha256_file(input_path),
        "source_type": ext[1:],
        "intermediate_docx": str(docx_path),
        "scanned_pdf": scanned,
        "warnings": warnings,
        "unit_count": len(units),
        "translatable_count": sum(1 for unit in units if unit["translatable"]),
        "units_file": "units.jsonl",
        "state_file": "state.json",
    }
    with (work / "units.jsonl").open("w", encoding="utf-8") as fh:
        for unit in units:
            fh.write(json.dumps(unit, ensure_ascii=False) + "\n")
    state = {
        "version": 1,
        "phase": "analysis",
        "analysis_done": [],
        "translation_done": [],
        "review_done": [],
        "terms": [],
        "glossary_locked": False,
        "translations": {},
        "review_corrections": {},
        "warnings": warnings,
    }
    write_json(work / "manifest.json", manifest)
    write_json(work / "state.json", state)
    print(json.dumps({"ok": True, "work_dir": str(work), "manifest": manifest}, ensure_ascii=False, indent=2))


def load_units(work: Path) -> list[dict[str, Any]]:
    path = work / "units.jsonl"
    if not path.exists():
        die(f"Missing prepared work directory: {work}")
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def load_state(work: Path) -> dict[str, Any]:
    state = read_json(work / "state.json")
    if state is None:
        die(f"Missing state.json in {work}")
    return state


def save_state(work: Path, state: dict[str, Any]) -> None:
    write_json(work / "state.json", state)


def choose_batch(units: list[dict[str, Any]], done: set[str], limit: int, chars: int) -> list[dict[str, Any]]:
    batch: list[dict[str, Any]] = []
    total = 0
    for unit in units:
        if unit["id"] in done or not unit["translatable"]:
            continue
        length = len(unit["text"])
        if batch and (len(batch) >= limit or total + length > chars):
            break
        batch.append(unit)
        total += length
    return batch


def next_batch(args: argparse.Namespace) -> None:
    work = Path(args.work_dir).expanduser().resolve()
    units = load_units(work)
    state = load_state(work)
    phase = args.phase
    if phase == "translation" and not state.get("glossary_locked"):
        die("Translation is locked until full analysis is complete and glossary.json is locked.")
    if phase == "review" and len(state.get("translations", {})) < sum(1 for u in units if u["translatable"]):
        die("Review is locked until every translatable unit has a translation.")
    done_key = {"analysis": "analysis_done", "translation": "translation_done", "review": "review_done"}[phase]
    batch = choose_batch(units, set(state.get(done_key, [])), args.limit, args.chars)
    terms = state.get("terms", [])
    payload = {
        "ok": True,
        "phase": phase,
        "batch_id": sha256_text(phase + "|" + "|".join(u["id"] for u in batch))[:16] if batch else None,
        "complete": not batch,
        "instructions": {
            "analysis": "Inspect every item, identify domain terminology, and return terms before any translation.",
            "translation": "Translate only the explicit item IDs. Preserve numbers, names, units, and intent.",
            "review": "Review each translation against source, glossary, and adjacent context; return corrections only when needed.",
        }[phase],
        "glossary": terms if phase != "analysis" else [],
        "items": [
            {
                "id": unit["id"],
                "source_hash": unit["source_hash"],
                "text": unit["text"],
                "normalized_text": unit["normalized_text"],
                "direction": unit["direction"],
                "kind": unit["kind"],
            }
            for unit in batch
        ],
    }
    write_json(work / f"current_{phase}_batch.json", payload)
    print(json.dumps(payload, ensure_ascii=False, indent=2))


def response_items(response: dict[str, Any]) -> list[dict[str, Any]]:
    items = response.get("items")
    if not isinstance(items, list):
        die("Response must contain an items array.")
    return items


def ingest(args: argparse.Namespace) -> None:
    work = Path(args.work_dir).expanduser().resolve()
    state = load_state(work)
    units = {unit["id"]: unit for unit in load_units(work)}
    response = read_json(Path(args.response))
    if not isinstance(response, dict):
        die("Response must be a JSON object.")
    phase = args.phase
    batch = read_json(work / f"current_{phase}_batch.json")
    if not batch:
        die(f"No current {phase} batch. Run next-batch first.")
    expected = {item["id"] for item in batch["items"]}
    received_items = response_items(response) if phase != "analysis" else response.get("items", [])
    received = {item.get("id") for item in received_items if isinstance(item, dict)}
    if phase != "analysis" and received != expected:
        die(f"Response IDs do not exactly match batch (missing={sorted(expected-received)}, extra={sorted(received-expected)}).")
    if phase == "analysis":
        terms = response.get("terms", [])
        if not isinstance(terms, list):
            die("Analysis response terms must be an array.")
        for term in terms:
            if not isinstance(term, dict) or not term.get("source") or not term.get("target"):
                continue
            state.setdefault("terms", []).append(
                {
                    "source": str(term["source"]).strip(),
                    "target": str(term["target"]).strip(),
                    "direction": term.get("direction", "auto"),
                    "domain": term.get("domain", ""),
                    "definition": term.get("definition", ""),
                    "context_unit_ids": term.get("context_unit_ids", []),
                    "confidence": term.get("confidence", "medium"),
                    "notes": term.get("notes", ""),
                }
            )
        state.setdefault("analysis_done", []).extend(sorted(expected))
    elif phase == "translation":
        for item in received_items:
            unit = units[item["id"]]
            if item.get("source_hash") != unit["source_hash"]:
                die(f"Source hash mismatch for {item['id']}.")
            target = item.get("target")
            if not isinstance(target, str) or not target.strip():
                die(f"Empty target for {item['id']}.")
            state.setdefault("translations", {})[item["id"]] = target
        state.setdefault("translation_done", []).extend(sorted(expected))
    else:
        for item in received_items:
            unit_id = item["id"]
            status = item.get("status", "ok")
            if status in {"replace", "correct", "fix"} and isinstance(item.get("target"), str) and item["target"].strip():
                state.setdefault("review_corrections", {})[unit_id] = item["target"].strip()
            state.setdefault("review_done", []).append(unit_id)
    state["phase"] = phase
    save_state(work, state)
    print(json.dumps({"ok": True, "phase": phase, "accepted": sorted(expected), "terms_total": len(state.get("terms", []))}, ensure_ascii=False, indent=2))


def dedupe_terms(terms: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    seen: dict[tuple[str, str], dict[str, Any]] = {}
    for term in terms:
        source, target = normalize_text(str(term.get("source", ""))), normalize_text(str(term.get("target", "")))
        if not source or not target:
            continue
        key = (source.casefold(), target.casefold())
        if key not in seen:
            item = dict(term)
            item["source"], item["target"] = source, target
            seen[key] = item
    return sorted(seen.values(), key=lambda x: x["source"].casefold())


def lock_glossary(args: argparse.Namespace) -> None:
    work = Path(args.work_dir).expanduser().resolve()
    state = load_state(work)
    units = load_units(work)
    expected = {u["id"] for u in units if u["translatable"]}
    if not expected.issubset(set(state.get("analysis_done", []))):
        die("Cannot lock glossary before every translatable unit has been analyzed.")
    supplied = read_json(Path(args.file))
    if isinstance(supplied, dict):
        supplied = supplied.get("terms", [])
    if not isinstance(supplied, list):
        die("Glossary must be a JSON array or an object with a terms array.")
    terms = dedupe_terms(supplied + state.get("terms", []))
    state["terms"] = terms
    state["glossary_locked"] = True
    state["phase"] = "translation"
    save_state(work, state)
    write_json(work / "glossary.json", {"version": 1, "terms": terms})
    csv_path = work / "glossary.csv"
    with csv_path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=["source", "target", "direction", "domain", "definition", "context_unit_ids", "confidence", "notes"])
        writer.writeheader()
        for term in terms:
            row = {key: term.get(key, "") for key in writer.fieldnames}
            if isinstance(row["context_unit_ids"], list):
                row["context_unit_ids"] = ",".join(map(str, row["context_unit_ids"]))
            writer.writerow(row)
    print(json.dumps({"ok": True, "glossary_locked": True, "terms": len(terms), "csv": str(csv_path)}, ensure_ascii=False, indent=2))


def find_first_run(p: Any) -> Any | None:
    runs = p.xpath("./w:r", namespaces=NS)
    return runs[0] if runs else None


def clear_paragraph_keep_properties(p: Any) -> None:
    for child in list(p):
        if child.tag != f"{W}pPr":
            p.remove(child)


def set_run_text(run: Any, text: str) -> None:
    for child in list(run):
        if child.tag != f"{W}rPr":
            run.remove(child)
    t = etree.SubElement(run, f"{W}t")
    if text[:1].isspace() or text[-1:].isspace():
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text


def set_visible_run_fonts(run: Any, text: str, translation: bool = False) -> None:
    """Use installed macOS fonts so both languages remain visible in DOCX/PDF."""
    rpr = run.find(f"{W}rPr")
    if rpr is None:
        rpr = etree.Element(f"{W}rPr")
        run.insert(0, rpr)
    fonts = rpr.find(f"{W}rFonts")
    if fonts is None:
        fonts = etree.Element(f"{W}rFonts")
        rpr.insert(0, fonts)
    if cjk_count(text):
        for name in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(f"{W}{name}", VISIBLE_CJK_FONT)
    elif translation:
        fonts.set(f"{W}ascii", VISIBLE_LATIN_FONT)
        fonts.set(f"{W}hAnsi", VISIBLE_LATIN_FONT)
        fonts.set(f"{W}eastAsia", VISIBLE_CJK_FONT)


def normalize_original_cjk_fonts(root: Any) -> None:
    for run in root.xpath(".//w:r", namespaces=NS):
        text = "".join(run.xpath(".//w:t/text()", namespaces=NS))
        if cjk_count(text):
            set_visible_run_fonts(run, text)


def make_translation_paragraph(original: Any, target: str) -> Any:
    translated = deepcopy(original)
    ppr = translated.find(f"{W}pPr")
    if ppr is None:
        ppr = etree.Element(f"{W}pPr")
        translated.insert(0, ppr)
    num_pr = ppr.find(f"{W}numPr")
    if num_pr is None:
        num_pr = etree.SubElement(ppr, f"{W}numPr")
    else:
        for child in list(num_pr):
            num_pr.remove(child)
    num_id = etree.SubElement(num_pr, f"{W}numId")
    num_id.set(f"{W}val", "0")
    clear_paragraph_keep_properties(translated)
    source_run = find_first_run(original)
    run = etree.SubElement(translated, f"{W}r")
    if source_run is not None:
        rpr = source_run.find(f"{W}rPr")
        if rpr is not None:
            run.append(deepcopy(rpr))
    set_run_text(run, target)
    set_visible_run_fonts(run, target, translation=True)
    return translated


def apply_translations_to_docx(source_docx: Path, output_docx: Path, units: list[dict[str, Any]], translations: dict[str, str]) -> dict[str, Any]:
    require_lxml()
    temp = Path(tempfile.mkdtemp(prefix="bilingual-docx-"))
    try:
        with zipfile.ZipFile(source_docx) as archive:
            archive.extractall(temp)
        inserted = 0
        by_part: dict[str, list[dict[str, Any]]] = {}
        for unit in units:
            if unit["id"] in translations:
                by_part.setdefault(unit["part"], []).append(unit)
        for part, part_units in by_part.items():
            path = temp / part
            if not path.exists():
                continue
            root = etree.parse(str(path), etree.XMLParser(remove_blank_text=False)).getroot()
            paragraphs = root.xpath(".//w:p", namespaces=NS)
            for unit in sorted(part_units, key=lambda u: u["paragraph_index"], reverse=True):
                index = unit["paragraph_index"]
                if index >= len(paragraphs):
                    continue
                original = paragraphs[index]
                if paragraph_text(original) != unit["text"]:
                    continue
                translated = make_translation_paragraph(original, translations[unit["id"]])
                parent = original.getparent()
                parent.insert(parent.index(original) + 1, translated)
                inserted += 1
            normalize_original_cjk_fonts(root)
            path.write_bytes(etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True))
        output_docx.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(output_docx, "w", zipfile.ZIP_DEFLATED) as archive:
            for path in sorted(temp.rglob("*")):
                if path.is_file():
                    archive.write(path, path.relative_to(temp).as_posix())
        return {"inserted_translations": inserted, "requested_translations": len(translations)}
    finally:
        shutil.rmtree(temp, ignore_errors=True)


def run_soffice_convert(docx_path: Path, output_dir: Path) -> tuple[bool, str]:
    candidates = [os.environ.get("SOFFICE"), shutil.which("soffice"), shutil.which("libreoffice")]
    command = next((item for item in candidates if item), None)
    if not command:
        return False, "LibreOffice/soffice not found; PDF export skipped."
    output_dir.mkdir(parents=True, exist_ok=True)
    profile = Path(tempfile.mkdtemp(prefix="soffice-profile-"))
    fontconfig = Path(tempfile.mkdtemp(prefix="soffice-fontconfig-"))
    try:
        cache = fontconfig / "cache"
        cache.mkdir()
        font_dirs = [
            Path.home() / "Library" / "Fonts",
            Path("/Library/Fonts"),
            Path("/System/Library/Fonts"),
            Path("/System/Library/AssetsV2"),
        ]
        directory_xml = "\n".join(f"  <dir>{path}</dir>" for path in font_dirs if path.exists())
        config_path = fontconfig / "fonts.conf"
        config_path.write_text(
            '<?xml version="1.0"?>\n<!DOCTYPE fontconfig SYSTEM "fonts.dtd">\n'
            f"<fontconfig>\n{directory_xml}\n  <cachedir>{cache}</cachedir>\n</fontconfig>\n",
            encoding="utf-8",
        )
        env = os.environ.copy()
        env["FONTCONFIG_FILE"] = str(config_path)
        env["XDG_CACHE_HOME"] = str(cache)
        result = subprocess.run(
            [command, "--headless", f"-env:UserInstallation={profile.as_uri()}", "--convert-to", "pdf", "--outdir", str(output_dir), str(docx_path)],
            capture_output=True,
            text=True,
            env=env,
        )
    finally:
        shutil.rmtree(profile, ignore_errors=True)
        shutil.rmtree(fontconfig, ignore_errors=True)
    if result.returncode != 0:
        return False, (result.stderr or result.stdout or "soffice conversion failed").strip()
    return True, (result.stdout or "PDF exported").strip()


def structural_counts(docx_path: Path) -> dict[str, int]:
    require_lxml()
    counts = {"media": 0, "tables": 0, "sections": 0}
    with zipfile.ZipFile(docx_path) as archive:
        counts["media"] = sum(name.startswith("word/media/") for name in archive.namelist())
        for part in list_story_parts(archive):
            root = etree.fromstring(archive.read(part))
            counts["tables"] += len(root.xpath(".//w:tbl", namespaces=NS))
            counts["sections"] += len(root.xpath(".//w:sectPr", namespaces=NS))
    return counts


def validate_output(source_docx: Path, output_docx: Path, units: list[dict[str, Any]], translations: dict[str, str]) -> dict[str, Any]:
    require_lxml()
    output_units = extract_docx_units(output_docx)
    texts = [unit["text"] for unit in output_units]
    cursor = 0
    missing_original: list[str] = []
    missing_translation: list[str] = []
    for unit in units:
        if not unit["translatable"]:
            continue
        try:
            pos = texts.index(unit["text"], cursor)
        except ValueError:
            missing_original.append(unit["id"])
            continue
        cursor = pos + 1
        target = translations.get(unit["id"], "")
        following = texts[cursor] if cursor < len(texts) else ""
        if normalize_text(following) != normalize_text(target):
            missing_translation.append(unit["id"])
        else:
            cursor += 1
    before, after = structural_counts(source_docx), structural_counts(output_docx)
    structural_ok = all(after[key] >= before[key] for key in before)
    return {
        "original_units": len(units),
        "output_units": len(output_units),
        "missing_original": missing_original,
        "missing_translation": missing_translation,
        "structure_before": before,
        "structure_after": after,
        "structure_ok": structural_ok,
        "passed": not missing_original and not missing_translation and structural_ok,
    }


def write_qa_report(path: Path, manifest: dict[str, Any], qa: dict[str, Any], render_message: str) -> None:
    lines = [
        "# Bilingual Translation QA Report",
        "",
        f"- Source type: `{manifest.get('source_type')}`",
        f"- Source SHA-256: `{manifest.get('input_sha256')}`",
        f"- Translatable units: `{manifest.get('translatable_count')}`",
        f"- Result: **{'PASS' if qa.get('passed') else 'REVIEW REQUIRED'}**",
        "",
        "## Structural checks",
        "",
        f"- Original units found: {qa.get('original_units')}",
        f"- Output units found: {qa.get('output_units')}",
        f"- Missing originals: {len(qa.get('missing_original', []))}",
        f"- Misplaced/missing translations: {len(qa.get('missing_translation', []))}",
        f"- Structure preserved: `{qa.get('structure_ok')}`",
        f"- PDF export: {render_message}",
        "",
        "## Warnings",
        "",
    ]
    warnings = manifest.get("warnings", [])
    lines.extend((f"- {warning}" for warning in warnings)) if warnings else lines.append("- None")
    if qa.get("missing_original"):
        lines.extend(["", "## Missing original IDs", "", *[f"- `{item}`" for item in qa["missing_original"]]])
    if qa.get("missing_translation"):
        lines.extend(["", "## Translation placement issues", "", *[f"- `{item}`" for item in qa["missing_translation"]]])
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def finalize(args: argparse.Namespace) -> None:
    work = Path(args.work_dir).expanduser().resolve()
    manifest = read_json(work / "manifest.json")
    stem = Path(manifest["input"]).stem
    output_root = Path(args.output_dir).expanduser().resolve() if args.output_dir else work / "output"
    output_dir = output_root if args.flat_output else output_root / stem
    output_dir.mkdir(parents=True, exist_ok=True)
    state = load_state(work)
    units = load_units(work)
    expected = {u["id"] for u in units if u["translatable"]}
    all_translations = dict(state.get("translations", {}))
    all_translations.update(state.get("review_corrections", {}))
    translations = {unit_id: all_translations[unit_id] for unit_id in expected if unit_id in all_translations}
    if not expected.issubset(translations):
        die("Cannot finalize: some translatable units are missing translations.")
    source_docx = Path(manifest["intermediate_docx"])
    output_docx = output_dir / f"{stem}.bilingual.docx"
    details = apply_translations_to_docx(source_docx, output_docx, units, translations)
    output_pdf = output_dir / f"{stem}.bilingual.pdf"
    if args.pdf:
        ok_pdf, pdf_message = run_soffice_convert(output_docx, output_dir)
    else:
        ok_pdf, pdf_message = False, "Skipped (DOCX-only fast mode)."
    qa = validate_output(source_docx, output_docx, units, translations)
    qa["insert_details"] = details
    qa["pdf_requested"] = bool(args.pdf)
    qa["pdf_exists"] = output_pdf.exists() if args.pdf and ok_pdf else False
    if ok_pdf and not output_pdf.exists():
        qa["passed"] = False
        pdf_message = "soffice returned success but the expected PDF was not found."
    glossary_source = work / "glossary.csv"
    glossary_output = output_dir / f"{stem}.glossary.csv"
    if glossary_source.exists():
        shutil.copy2(glossary_source, glossary_output)
    qa_path = output_dir / f"{stem}.qa-report.md"
    qa_json_path = output_dir / f"{stem}.qa.json"
    write_qa_report(qa_path, manifest, qa, pdf_message)
    write_json(qa_json_path, qa)
    state["phase"] = "complete" if qa["passed"] else "review-required"
    state["qa"] = qa
    save_state(work, state)
    outputs = [str(output_docx), str(glossary_output), str(qa_path), str(qa_json_path)]
    if args.pdf:
        outputs.insert(1, str(output_pdf))
    print(
        json.dumps(
            {
                "ok": qa["passed"],
                "output_dir": str(output_dir),
                "qa": qa,
                "outputs": outputs,
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    if not qa["passed"]:
        raise SystemExit(1)


def validate_command(args: argparse.Namespace) -> None:
    work = Path(args.work_dir).expanduser().resolve()
    manifest = read_json(work / "manifest.json")
    state = load_state(work)
    units = load_units(work)
    expected = {u["id"] for u in units if u["translatable"]}
    result = {
        "analysis_complete": expected.issubset(set(state.get("analysis_done", []))),
        "glossary_locked": bool(state.get("glossary_locked")),
        "translations_complete": expected.issubset(set(state.get("translations", {}))),
        "review_complete": expected.issubset(set(state.get("review_done", []))),
        "manifest": manifest,
    }
    result["passed"] = all(result[key] for key in ("analysis_complete", "glossary_locked", "translations_complete", "review_complete"))
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if not result["passed"]:
        raise SystemExit(1)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("prepare")
    p.add_argument("input")
    p.add_argument("--work-dir", required=True)
    p.set_defaults(func=prepare)
    p = sub.add_parser("next-batch")
    p.add_argument("--work-dir", required=True)
    p.add_argument("--phase", choices=["analysis", "translation", "review"], required=True)
    p.add_argument("--limit", type=int, default=8)
    p.add_argument("--chars", type=int, default=9000)
    p.set_defaults(func=next_batch)
    p = sub.add_parser("ingest")
    p.add_argument("--work-dir", required=True)
    p.add_argument("--phase", choices=["analysis", "translation", "review"], required=True)
    p.add_argument("--response", required=True)
    p.set_defaults(func=ingest)
    p = sub.add_parser("lock-glossary")
    p.add_argument("--work-dir", required=True)
    p.add_argument("--file", required=True)
    p.set_defaults(func=lock_glossary)
    p = sub.add_parser("finalize")
    p.add_argument("--work-dir", required=True)
    p.add_argument(
        "--output-dir",
        help="Output root; creates <output-dir>/<source-filename>/ by default",
    )
    p.add_argument(
        "--flat-output",
        action="store_true",
        help="Write directly into --output-dir instead of creating a source-named folder",
    )
    p.add_argument("--pdf", action="store_true", help="Also export PDF; skipped by default")
    p.set_defaults(func=finalize)
    p = sub.add_parser("validate")
    p.add_argument("--work-dir", required=True)
    p.set_defaults(func=validate_command)
    return parser


if __name__ == "__main__":
    args = build_parser().parse_args()
    args.func(args)
