from __future__ import annotations

import argparse
import json
import sys
import tempfile
import unittest
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from document_pipeline import (  # noqa: E402
    docx_format_signature,
    extract_docx_units,
    finalize,
    prepare,
    read_json,
    structural_counts,
    write_json,
)


class EnglishCopyTests(unittest.TestCase):
    def test_chinese_docx_adds_translation_only_english_copy(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            source = temp / "设备清单.docx"
            work = temp / "work"
            output_root = temp / "Output Files"

            document = Document()
            document.add_heading("设备清单", level=1)
            paragraph = document.add_paragraph()
            paragraph.add_run("设备").bold = True
            paragraph.add_run("说明")
            document.add_paragraph("这是用于医院设备维护和安全操作的详细说明。")
            document.add_paragraph("术语 Technical Term")
            document.add_paragraph("Safety Instructions")
            document.add_paragraph("MODEL ABC-123")
            table = document.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "备用零件"
            document.save(source)

            prepare(argparse.Namespace(input=str(source), work_dir=str(work)))
            manifest = read_json(work / "manifest.json")
            self.assertEqual(manifest["document_direction"], "zh-to-en")
            self.assertTrue(manifest["english_copy_required"])
            units = [
                json.loads(line)
                for line in (work / "units.jsonl").read_text(encoding="utf-8").splitlines()
            ]
            targets = {
                "设备清单": "Equipment List",
                "设备说明": "Equipment Description",
                "这是用于医院设备维护和安全操作的详细说明。": "This provides detailed guidance for hospital equipment maintenance and safe operation.",
                "术语 Technical Term": "Technical Term",
                "MODEL ABC-123": "型号 ABC-123",
                "备用零件": "Spare Parts",
            }
            english_unit = next(
                unit for unit in units if unit["text"] == "Safety Instructions"
            )
            self.assertFalse(english_unit["translatable"])
            self.assertFalse(english_unit["bilingual_output"])
            self.assertEqual(
                english_unit["translation_reason"],
                "preserve_existing_english_once",
            )
            translations = {
                unit["id"]: targets[unit["text"]]
                for unit in units
                if unit["translatable"]
            }
            state = read_json(work / "state.json")
            state.update(
                {
                    "phase": "review",
                    "analysis_done": list(translations),
                    "translation_done": list(translations),
                    "review_done": list(translations),
                    "glossary_locked": True,
                    "translations": translations,
                }
            )
            write_json(work / "state.json", state)

            finalize(
                argparse.Namespace(
                    work_dir=str(work),
                    output_dir=str(output_root),
                    flat_output=False,
                    pdf=False,
                )
            )

            output_dir = output_root / "设备清单"
            bilingual = output_dir / "设备清单.bilingual.docx"
            english = output_dir / "设备清单.english.docx"
            self.assertTrue(bilingual.is_file())
            self.assertTrue(english.is_file())
            self.assertEqual(structural_counts(source), structural_counts(english))
            self.assertEqual(docx_format_signature(source), docx_format_signature(english))

            source_ids = {unit["text"]: unit["id"] for unit in units}
            english_map = {unit["id"]: unit["text"] for unit in extract_docx_units(english)}
            for source_text in (
                "设备清单",
                "设备说明",
                "这是用于医院设备维护和安全操作的详细说明。",
                "术语 Technical Term",
                "备用零件",
            ):
                self.assertEqual(
                    english_map[source_ids[source_text]],
                    targets[source_text],
                )
            self.assertEqual(english_map[source_ids["MODEL ABC-123"]], "MODEL ABC-123")
            self.assertEqual(
                english_map[source_ids["Safety Instructions"]],
                "Safety Instructions",
            )

            bilingual_texts = [unit["text"] for unit in extract_docx_units(bilingual)]
            self.assertIn("术语 Technical Term", bilingual_texts)
            self.assertEqual(bilingual_texts.count("Technical Term"), 0)
            self.assertEqual(bilingual_texts.count("Safety Instructions"), 1)

            qa = read_json(output_dir / "设备清单.qa.json")
            self.assertTrue(qa["passed"])
            self.assertTrue(qa["english_copy"]["required"])
            self.assertTrue(qa["english_copy"]["passed"])
            self.assertTrue(qa["english_copy"]["format_ok"])
            self.assertEqual(qa["english_copy"]["residual_chinese_targets"], [])


if __name__ == "__main__":
    unittest.main()
